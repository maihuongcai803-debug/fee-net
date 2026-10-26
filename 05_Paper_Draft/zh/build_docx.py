# -*- coding: utf-8 -*-
"""Render the Chinese manuscript to an IEEE-formatted .docx."""
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, "05_Paper_Draft/zh")
import content as C  # noqa: E402

OUT = "05_Paper_Draft/FEE-Net_中文稿_EI会议格式.docx"

LATIN = "Times New Roman"
SONG = "宋体"
HEI = "黑体"

# ---------------------------------------------------------------- numbering
order, seen = [], set()
for blk in C.BLOCKS:
    texts = []
    if blk[0] in ("para", "subsection", "section"):
        texts = [blk[1]]
    elif blk[0] == "enum":
        texts = list(blk[1])
    elif blk[0] == "table":
        texts = [blk[1]] + [c for r in blk[3] for c in r]
    elif blk[0] == "figure":
        texts = [blk[2]]
    for t in texts:
        for m in re.finditer(r"\[\[([^\]]+)\]\]", t):
            for k in (x.strip() for x in m.group(1).split(",")):
                if k not in seen:
                    seen.add(k)
                    order.append(k)
NUM = {k: i + 1 for i, k in enumerate(order)}


def cite(m):
    nums = sorted(NUM[k] for k in (x.strip() for x in m.group(1).split(",")))
    out, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        out.append(str(nums[i]) if j == i else
                   (f"{nums[i]}, {nums[j]}" if j == i + 1 else f"{nums[i]}–{nums[j]}"))
        i = j + 1
    return "[" + ", ".join(out) + "]"


def resolve(t):
    return re.sub(r"\[\[([^\]]+)\]\]", cite, t)


# ---------------------------------------------------------------- primitives
def font(run, size, bold=False, italic=False, latin=LATIN, east=SONG):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.name = latin
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), latin)
    rFonts.set(qn("w:hAnsi"), latin)
    rFonts.set(qn("w:eastAsia"), east)


EQ_PAT = re.compile(r"_\{([^}]*)\}|\^\{([^}]*)\}")


def add_inline(p, text, size, bold=False, italic=False, east=SONG):
    """Add resolved text to a paragraph, honouring _{...} / ^{...} markers."""
    text = resolve(text)
    pos = 0
    for m in EQ_PAT.finditer(text):
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()])
            font(r, size, bold, italic, east=east)
        seg = m.group(1) if m.group(1) is not None else (m.group(2) or "")
        r = p.add_run(seg)
        font(r, size, bold, italic, east=east)
        if m.group(1) is not None:
            r.font.subscript = True
        else:
            r.font.superscript = True
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:])
        font(r, size, bold, italic, east=east)


def para(doc, text, *, size=10.5, bold=False, italic=False, align="justify",
         indent=True, space_before=0, space_after=0, line=1.0, east=SONG,
         keep_with_next=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                    "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if indent:
        pf.first_line_indent = Cm(size * 0.074)  # ~2 Chinese characters
    pf.keep_with_next = keep_with_next
    add_inline(p, text, size, bold, italic, east)
    return p


def rich_para(doc, parts, *, size=9, align="justify", indent=False,
              space_before=0, space_after=0):
    """parts: list of (text, bold, italic)."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                    "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.0
    if indent:
        pf.first_line_indent = Cm(0.6)
    for text, b, i in parts:
        r = p.add_run(resolve(text))
        font(r, size, b, i)
    return p


def add_eq_runs(p, text, size=10.5):
    pos = 0
    for m in EQ_PAT.finditer(text):
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()])
            font(r, size, italic=False)
        seg = m.group(1) or m.group(2) or ""
        r = p.add_run(seg)
        font(r, size)
        if m.group(1):
            r.font.subscript = True
        else:
            r.font.superscript = True
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:])
        font(r, size)


def eq_para(doc, text, number, usable_cm, size=10.5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(4)
    pf.space_after = Pt(4)
    pf.tab_stops.add_tab_stop(Cm(usable_cm / 2), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Cm(usable_cm), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    add_eq_runs(p, text, size)
    p.add_run("\t(" + number + ")")
    return p


def set_cell(cell, text, size=8.5, bold=False, align="center", east=SONG):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.alignment = (WD_ALIGN_PARAGRAPH.CENTER if align == "center"
                                    else WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.space_before = Pt(1.5)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.0
    add_inline(p, text, size, bold, False, east)
    # vertical centring
    tcPr = cell._tc.get_or_add_tcPr()
    va = OxmlElement("w:vAlign")
    va.set(qn("w:val"), "center")
    tcPr.append(va)


def cell_border(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    el = borders.find(qn("w:" + edge))
    if el is None:
        el = OxmlElement("w:" + edge)
        borders.append(el)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:color"), "000000")


def three_line_table(doc, headers, rows, widths, total_cm):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.style = "Table Grid"
    # remove all borders first
    tbl = t._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "none")
        borders.append(el)
    tblPr.append(borders)

    scale = total_cm / sum(widths)
    for r_i, row in enumerate([headers] + rows):
        for c_i, val in enumerate(row):
            cell = t.cell(r_i, c_i)
            cell.width = Cm(widths[c_i] * scale)
            set_cell(cell, val, bold=(r_i == 0))
    # keep the table on one page: no row splitting, and keep-with-next on all but the last row
    for r_i, row in enumerate(t.rows):
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
        if r_i < len(t.rows) - 1:
            for cell in row.cells:
                for para_ in cell.paragraphs:
                    para_.paragraph_format.keep_with_next = True
    # three lines
    for c in range(len(headers)):
        cell_border(t.cell(0, c), "top", 12)
        cell_border(t.cell(0, c), "bottom", 6)
        cell_border(t.cell(len(rows), c), "bottom", 12)
    return t


# ---------------------------------------------------------------- document
doc = Document()
st = doc.styles["Normal"]
st.font.name = LATIN
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(2.4)
sec.left_margin, sec.right_margin = Cm(1.6), Cm(1.6)
USABLE = 21.0 - 3.2

# conference papers carry no running head; leave the header empty


# title
tp = doc.add_paragraph()
tp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
tp.paragraph_format.space_before = Pt(0)
tp.paragraph_format.space_after = Pt(8)
tr = tp.add_run(C.TITLE)
font(tr, 18, bold=False, east=HEI)

ap = doc.add_paragraph()
ap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
ap.paragraph_format.space_after = Pt(4)
ar = ap.add_run(C.AUTHORS)
font(ar, 11, east=HEI)

for i, line in enumerate(C.AFFIL_LINES):
    fp = doc.add_paragraph()
    fp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_after = Pt(0 if i < len(C.AFFIL_LINES) - 1 else 10)
    fr = fp.add_run(line)
    font(fr, 9)

# abstract & keywords
rich_para(doc, [("摘要—", True, True), (C.ABSTRACT, False, False)], size=9,
          space_before=2, space_after=6)
rich_para(doc, [("关键词—", True, True), (C.KEYWORDS, False, False)], size=9,
          space_before=0, space_after=10)

eq_no = 0
eq_ptr = 0
for blk in C.BLOCKS:
    kind = blk[0]
    if kind == "section":
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        add_inline(p, blk[1], 10.5, True, False, HEI)
    elif kind == "subsection":
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        add_inline(p, blk[1], 10.5, True, True, HEI)
    elif kind == "para":
        para(doc, blk[1], space_before=0, space_after=0)
    elif kind == "enum":
        for i, item in enumerate(blk[1], 1):
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.left_indent = Cm(0.74)
            pf.first_line_indent = Cm(-0.37)
            pf.space_after = Pt(0)
            add_inline(p, f"{i}) " + item, 10.5)
    elif kind == "eq":
        tex, variants, nums = blk[1], blk[2], blk[3]
        for v, n in zip(variants, nums):
            eq_para(doc, v, n, USABLE)
    elif kind == "table":
        cap, headers, rows, widths = blk[1], blk[2], blk[3], blk[4]
        cp = doc.add_paragraph()
        cp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(8)
        cp.paragraph_format.space_after = Pt(3)
        cp.paragraph_format.keep_with_next = True
        add_inline(cp, cap, 8.5)
        three_line_table(doc, headers, rows, widths, USABLE)
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(6)
        sp.paragraph_format.line_spacing = 0.6
    elif kind == "figure":
        path, cap, width = blk[1], blk[2], blk[3]
        ip = doc.add_paragraph()
        ip.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ip.paragraph_format.space_before = Pt(8)
        ip.paragraph_format.keep_with_next = True
        ip.add_run().add_picture("05_Paper_Draft/" + path, width=Cm(min(width, USABLE)))
        cp = doc.add_paragraph()
        cp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(8)
        add_inline(cp, cap, 8.5)

# references
rp = doc.add_paragraph()
rp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
rp.paragraph_format.space_before = Pt(12)
rp.paragraph_format.space_after = Pt(6)
rr = rp.add_run("参考文献")
font(rr, 10.5, bold=True, east=HEI)

for k in order:
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.left_indent = Cm(0.75)
    pf.first_line_indent = Cm(-0.75)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.0
    r = p.add_run(f"[{NUM[k]}]  {C.REFS[k]}")
    font(r, 8.5)

doc.save(OUT)
print("wrote", OUT, "| references:", len(order), "| equations:", sum(len(b[3]) for b in C.BLOCKS if b[0] == "eq"))
