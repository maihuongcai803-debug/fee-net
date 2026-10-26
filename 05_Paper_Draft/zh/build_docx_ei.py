# -*- coding: utf-8 -*-
"""Render the Chinese EI-conference manuscript (.docx).

Layout follows the Chinese EI-conference norm:
  A4, 2.5 cm margins, two-column body, 10.5 pt Song body text,
  Arabic section numbering (1 / 1.1), bilingual title-abstract block,
  CLC + document code, bilingual table/figure captions, GB/T 7714 references.
Wide tables and the six-panel figure are placed in one-column sections.
"""
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, "05_Paper_Draft/zh")
import content as C          # noqa: E402
import meta_cn as M          # noqa: E402
from refs_gbt import REFS_GBT  # noqa: E402

OUT = "05_Paper_Draft/FEE-Net_中文稿_EI会议格式.docx"

SONG, HEI, KAI, TNR = "宋体", "黑体", "楷体", "Times New Roman"
PAGE_W, MARGIN = 21.0, 2.5
USABLE = PAGE_W - 2 * MARGIN          # 16.0 cm
COL_W = (USABLE - 0.8) / 2            # ≈7.6 cm per column

# ---------------------------------------------------------------- citations
order, seen = [], set()
for blk in C.BLOCKS:
    texts = []
    if blk[0] in ("para", "subsection", "section"):
        texts = [blk[1]]
    elif blk[0] == "enum":
        texts = list(blk[1])
    elif blk[0] == "table":
        texts = [blk[1]] + [c for r in blk[3] for c in r]
    elif blk[0] == "figure":
        texts = [blk[2]]
    for t in texts:
        for m in re.finditer(r"\[\[([^\]]+)\]\]", t):
            for k in (x.strip() for x in m.group(1).split(",")):
                if k not in seen:
                    seen.add(k)
                    order.append(k)
NUM = {k: i + 1 for i, k in enumerate(order)}


def cite(m):
    nums = sorted(NUM[k] for k in (x.strip() for x in m.group(1).split(",")))
    out, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        out.append(str(nums[i]) if j == i else
                   (f"{nums[i]},{nums[j]}" if j == i + 1 else f"{nums[i]}-{nums[j]}"))
        i = j + 1
    return "[" + ",".join(out) + "]"


def resolve(t):
    return re.sub(r"\[\[([^\]]+)\]\]", cite, t)


EQ = re.compile(r"_\{([^}]*)\}|\^\{([^}]*)\}")


def font(run, size, bold=False, italic=False, latin=TNR, east=SONG):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.name = latin
    rPr = run._element.get_or_add_rPr()
    rF = rPr.find(qn("w:rFonts"))
    if rF is None:
        rF = OxmlElement("w:rFonts")
        rPr.append(rF)
    for a in ("w:ascii", "w:hAnsi"):
        rF.set(qn(a), latin)
    rF.set(qn("w:eastAsia"), east)


def add_inline(p, text, size, bold=False, italic=False, east=SONG, latin=TNR):
    text = resolve(text)
    pos = 0
    for m in EQ.finditer(text):
        if m.start() > pos:
            font(p.add_run(text[pos:m.start()]), size, bold, italic, latin, east)
        r = p.add_run(m.group(1) if m.group(1) is not None else (m.group(2) or ""))
        font(r, size, bold, italic, latin, east)
        if m.group(1) is not None:
            r.font.subscript = True
        else:
            r.font.superscript = True
        pos = m.end()
    if pos < len(text):
        font(p.add_run(text[pos:]), size, bold, italic, latin, east)


def para(doc, text="", *, size=10.5, bold=False, italic=False, align="justify",
         indent=True, before=0, after=0, east=SONG, latin=TNR, line=1.0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                    "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = line
    if indent:
        pf.first_line_indent = Pt(size * 2)      # 首行缩进两字符
    if text:
        add_inline(p, text, size, bold, italic, east, latin)
    return p


def rich(doc, parts, *, size=9, align="justify", before=0, after=0, east=SONG, latin=TNR):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
                    "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = 1.0
    for text, b, it, ea, la, sz in parts:
        add_inline(p, text, sz or size, b, it, ea or east, la or latin)
    return p


def set_cols(section, num, space=0.8):
    sectPr = section._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sectPr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(int(space / 2.54 * 1440)))


def cell_border(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    bs = tcPr.find(qn("w:tcBorders"))
    if bs is None:
        bs = OxmlElement("w:tcBorders")
        tcPr.append(bs)
    el = bs.find(qn("w:" + edge))
    if el is None:
        el = OxmlElement("w:" + edge)
        bs.append(el)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:color"), "000000")


def three_line(doc, headers, rows, total_cm):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblPr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "none")
        borders.append(el)
    tblPr.append(borders)

    per = total_cm / len(headers)
    for r_i, row in enumerate([headers] + rows):
        for c_i, val in enumerate(row):
            cell = t.cell(r_i, c_i)
            cell.width = Cm(per)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            add_inline(p, val, 8.5 if r_i == 0 else 8, bold=(r_i == 0))
            tcPr = cell._tc.get_or_add_tcPr()
            va = OxmlElement("w:vAlign")
            va.set(qn("w:val"), "center")
            tcPr.append(va)
    for r_i, row in enumerate(t.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        if r_i < len(t.rows) - 1:
            for cell in row.cells:
                for pp in cell.paragraphs:
                    pp.paragraph_format.keep_with_next = True
    for c in range(len(headers)):
        cell_border(t.cell(0, c), "top", 12)
        cell_border(t.cell(0, c), "bottom", 6)
        cell_border(t.cell(len(rows), c), "bottom", 12)
    return t


def caption(doc, cn, en, above=True, width_note=None):
    """Bilingual caption: Chinese line then English line."""
    for text, lang in ((cn, "cn"), (en, "en")):
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3 if lang == "cn" else 0)
        p.paragraph_format.space_after = Pt(0 if lang == "cn" else 3)
        p.paragraph_format.keep_with_next = above
        p.paragraph_format.line_spacing = 1.0
        add_inline(p, text, 9, bold=(lang == "cn"),
                   east=HEI, latin=TNR if lang == "en" else TNR)


# ---------------------------------------------------------------- document
doc = Document()
st = doc.styles["Normal"]
st.font.name = TNR
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(PAGE_W), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(MARGIN)
sec.left_margin = sec.right_margin = Cm(MARGIN)
set_cols(sec, 1)

# ---------------- front matter (single column) ----------------
p = para(doc, C.TITLE, size=18, align="center", indent=False, before=0, after=6, east=HEI)
p = para(doc, C.AUTHORS, size=12, align="center", indent=False, after=1, east=KAI)
for line in C.AFFIL_LINES:
    para(doc, line, size=9, align="center", indent=False, after=0, east=SONG)
para(doc, "", size=9, indent=False, after=0)

rich(doc, [("摘要：", True, False, HEI, TNR, 9), (C.ABSTRACT, False, False, SONG, TNR, 9)],
     before=4, after=2)
rich(doc, [("关键词：", True, False, HEI, TNR, 9), (C.KEYWORDS, False, False, SONG, TNR, 9)],
     after=2)
rich(doc, [("中图分类号：", True, False, HEI, TNR, 9), (M.CLC, False, False, SONG, TNR, 9),
           ("　　文献标识码：", True, False, HEI, TNR, 9), (M.DOC_CODE, False, False, SONG, TNR, 9),
           ("　　文章编号：", True, False, HEI, TNR, 9), ("（由编辑部填写）", False, False, SONG, TNR, 9)],
     after=2)
rich(doc, [(M.FUND, False, False, SONG, TNR, 7.5)], after=0)
rich(doc, [(M.BIO, False, False, SONG, TNR, 7.5)], after=8)

rich(doc, [(M.TITLE_EN, True, False, TNR, TNR, 14)], align="center", after=3)
rich(doc, [(M.AUTHORS_EN, False, False, TNR, TNR, 10.5)], align="center", after=1)
for line in M.AFFIL_EN:
    rich(doc, [(line, False, False, TNR, TNR, 9)], align="center", after=0)
rich(doc, [("", False, False, TNR, TNR, 9)], after=2)
rich(doc, [("Abstract: ", True, False, TNR, TNR, 9), (M.ABSTRACT_EN, False, False, TNR, TNR, 9)],
     before=2, after=2)
rich(doc, [("Key words: ", True, False, TNR, TNR, 9), (M.KEYWORDS_EN, False, False, TNR, TNR, 9)],
     after=6)

# ---------------- body (two columns) ----------------
body = doc.add_section(WD_SECTION.CONTINUOUS)
body.top_margin = body.bottom_margin = Cm(MARGIN)
body.left_margin = body.right_margin = Cm(MARGIN)
set_cols(body, 2)

WIDE_TABLES = {2, 3, 5}     # 需要通栏的表序号
full_sections = []


def switch_to(num, width_is_full=None):
    s = doc.add_section(WD_SECTION.CONTINUOUS)
    s.top_margin = s.bottom_margin = Cm(MARGIN)
    s.left_margin = s.right_margin = Cm(MARGIN)
    set_cols(s, num)
    return s


tab_i = fig_i = 0
for blk in C.BLOCKS:
    kind = blk[0]
    if kind == "section":
        label = blk[1]
        m = re.match(r"^(\d)\s+(.*)$", label)
        txt = f"{m.group(1)}　{m.group(2)}" if m else label
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        add_inline(p, txt, 12, bold=True, east=HEI)
    elif kind == "subsection":
        m = re.match(r"^(\d\.\d)\s+(.*)$", blk[1])
        txt = f"{m.group(1)}　{m.group(2)}" if m else blk[1]
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        add_inline(p, txt, 10.5, bold=True, east=HEI)
    elif kind == "para":
        para(doc, blk[1], before=0, after=0)
    elif kind == "enum":
        for i, item in enumerate(blk[1], 1):
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.first_line_indent = Pt(21)
            pf.space_after = Pt(0)
            add_inline(p, f"（{i}）{item}", 10.5)
    elif kind == "eq":
        for v, n in zip(blk[2], blk[3]):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, f"{v}　　（{n}）", 10.5)
    elif kind == "table":
        cap, headers, rows = blk[1], blk[2], blk[3]
        tab_i += 1
        wide = tab_i in WIDE_TABLES
        if wide:
            switch_to(1, True)
            width = USABLE
        else:
            width = COL_W
        cn = re.sub(r"^(表\d+)\s*", "", cap)
        en = M.CAPTION_EN.get(f"表{tab_i}", "")
        caption(doc, f"表{tab_i}　{cn}", f"Table {tab_i}　{en}", above=True)
        three_line(doc, headers, rows, width)
        para(doc, "", size=4, indent=False, after=2)
        if wide:
            switch_to(2, False)
    elif kind == "figure":
        path, cap, width = blk[1], blk[2], blk[3]
        fig_i += 1
        wide = width > 9.0
        if wide:
            switch_to(1, True)
            img_w = USABLE
        else:
            img_w = COL_W
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture("05_Paper_Draft/" + path, width=Cm(img_w))
        cn = re.sub(r"^(图\d+)\s*", "", cap)
        en = M.CAPTION_EN.get(f"图{fig_i}", "")
        caption(doc, f"图{fig_i}　{cn}", f"Fig.{fig_i}　{en}", above=False)
        if wide:
            switch_to(2, False)

# ---------------- references ----------------
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(3)
p.paragraph_format.keep_with_next = True
add_inline(p, "参考文献", 12, bold=True, east=HEI)

for k in order:
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.left_indent = Cm(0.7)
    pf.first_line_indent = Cm(-0.7)
    pf.space_after = Pt(1)
    pf.line_spacing = 1.0
    add_inline(p, f"[{NUM[k]}]　{REFS_GBT[k]}", 8.5)

try:
    doc.save(OUT)
except PermissionError:
    alt = OUT.replace(".docx", "_v2.docx")
    doc.save(alt)
    print("[warn] 目标文件被占用，已另存为", alt)
    OUT = alt
print("wrote", OUT, "| refs", len(order), "| tables", tab_i, "| figures", fig_i,
      "| eqs", sum(len(b[3]) for b in C.BLOCKS if b[0] == "eq"))
